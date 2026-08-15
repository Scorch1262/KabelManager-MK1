# -*- coding: utf-8 -*-
"""
netlist_export.py
==================
Exportiert eine "Netzliste" (Tabelle aller Einzeladern) als Excel-Datei:
Von Geraet/Anschluss -> Nach Geraet/Anschluss, inkl. Kabelfarbe.
Standard-Verkabelungen (Buendel) werden dabei in ihre einzelnen Adern
aufgeloest (eine Zeile je Ader).
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def export_netlist_xlsx(path, netz_config):
    elements = netz_config.get("elements", [])
    connections = netz_config.get("connections", [])
    by_id = {el["id"]: el for el in elements}

    wb = Workbook()
    ws = wb.active
    ws.title = "Netzliste"

    headers = ["Nr.", "Von Geraet", "Von Standort", "Von Anschluss",
               "Nach Geraet", "Nach Standort", "Nach Anschluss",
               "Kabelfarbe", "Typ", "Bezeichnung / Signal"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="333B47")
        cell.alignment = Alignment(horizontal="center")

    row_num = 0
    for conn in connections:
        fel = by_id.get(conn["from_element"])
        tel = by_id.get(conn["to_element"])
        fname = fel["name"] if fel else "?"
        ford = fel.get("location", "") if fel else ""
        tname = tel["name"] if tel else "?"
        tord = tel.get("location", "") if tel else ""

        if conn["kind"] == "einzel":
            row_num += 1
            ws.append([
                row_num, fname, ford, conn.get("from_pin") or "-",
                tname, tord, conn.get("to_pin") or "-",
                conn.get("color", ""), "Einzelkabel", conn.get("label", ""),
            ])
        else:
            for wire in conn.get("wires", []):
                row_num += 1
                von_anschluss = wire.get("von_anschluss") or (conn.get("from_pin") or "-")
                nach_anschluss = wire.get("nach_anschluss") or (conn.get("to_pin") or "-")
                ws.append([
                    row_num, fname, ford, von_anschluss,
                    tname, tord, nach_anschluss,
                    wire.get("farbe", ""), conn.get("standard", ""), wire.get("signal", ""),
                ])

    widths = [6, 22, 16, 20, 22, 16, 20, 12, 20, 20]
    for i, w in enumerate(widths, start=1):
        col_letter = chr(64 + i)
        ws.column_dimensions[col_letter].width = w

    ws.freeze_panes = "A2"
    wb.save(path)
